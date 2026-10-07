from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


EMAIL_RE = re.compile(r"[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}")
PHONE_RE = re.compile(r"(?<!\d)(?:0\d[\d\s\-]{7,}\d|\d{4,5})(?!\d)")
BUILDING_RE = re.compile(r"(?:ตึก|อาคาร)\s*([0-9]{3})", re.IGNORECASE)
FLOOR_RE = re.compile(r"ชั้น(?:ที่)?\s*([0-9]{1,2}|G|M)\b", re.IGNORECASE)
HONORIFICS = ("คุณ", "นาย", "นาง", "นางสาว", "ดร.", "ดร ")


@dataclass
class ParsedRow:
    row_number: int
    classification: str
    raw_values: list[str]
    parsed_data: dict[str, Any] | None = None
    issues: list[tuple[str, str, str]] = field(default_factory=list)


@dataclass
class ParsedSheet:
    name: str
    detected_building: str | None
    detected_floor: str | None
    rows: list[ParsedRow]

    @property
    def nonempty_rows(self) -> int:
        return len(self.rows)

    @property
    def employee_rows(self) -> int:
        return sum(row.classification in {"employee", "employee_review"} for row in self.rows)

    @property
    def issue_rows(self) -> int:
        return sum(bool(row.issues) for row in self.rows)


def clean_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def normalize_text(value: str | None) -> str:
    return re.sub(r"\s+", "", (value or "").strip().lower())


def digits_only(value: str | None) -> str:
    return re.sub(r"\D", "", value or "")


def _canonical_column(label: str) -> str | None:
    text = normalize_text(label).replace("-", "")
    if not text:
        return None
    if "ชื่อเล่น" in text:
        return "nickname"
    if text == "ชื่อ" or text.startswith("ชื่อ–") or text.startswith("ชื่อ-"):
        return "full_name"
    if "ตำแหน่ง" in text:
        return "position"
    if "email" in text:
        return "email"
    if "เบอร์ต่อ" in text or "ต่อภายใน" in text or text == "ต่อ":
        return "extension"
    if "เบอร์โทร" in text or "โทรศัพท์" in text or "โทรติดต่อ" in text:
        return "phone"
    if any(word in text for word in ("หน่วยงาน", "สังกัด", "แผนก", "ฝ่าย")):
        return "organization"
    return None


def _header_map(values: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for index, value in enumerate(values):
        key = _canonical_column(value)
        if key and key not in mapping:
            mapping[key] = index
    if "full_name" in mapping and len(mapping) >= 3:
        return mapping
    return None


def _find_building(file_name: str, sheet_name: str, leading_text: str) -> str | None:
    for source in (sheet_name, leading_text, file_name):
        match = BUILDING_RE.search(source)
        if match:
            return match.group(1)
    return None


def _find_floor(sheet_name: str, leading_text: str) -> str | None:
    for source in (sheet_name, leading_text):
        match = FLOOR_RE.search(source)
        if match:
            return match.group(1).upper()
    return None


def _looks_like_person_name(value: str) -> bool:
    compact = value.strip()
    return any(compact.startswith(prefix) for prefix in HONORIFICS) or bool(re.search(r"[ก-๙]{2,}\s+[ก-๙]{2,}", compact))


def _infer_unit_type(name: str | None) -> str:
    text = name or ""
    if "แผนก" in text:
        return "แผนก"
    if "ฝ่าย" in text:
        return "ฝ่าย"
    if "สำนัก" in text:
        return "สำนัก"
    if "งาน" in text:
        return "งาน"
    return "หน่วยงาน"


def _extract_free_text(text: str, unit: str | None) -> tuple[dict[str, Any] | None, list[tuple[str, str, str]]]:
    emails = EMAIL_RE.findall(text)
    phones = [p.strip() for p in PHONE_RE.findall(text)]
    if not any(prefix in text for prefix in HONORIFICS) or not (emails or phones):
        return None, []
    cut_points = [len(text)]
    for marker in (" โทร", " เบอร์", " E-mail", " email", " แม่บ้าน", " พนักงาน", " คนขับ"):
        at = text.lower().find(marker.lower())
        if at > 0:
            cut_points.append(at)
    name = text[: min(cut_points)].strip(" :-")
    parsed = {
        "full_name": name,
        "nickname": None,
        "position": None,
        "email": emails[0].lower() if emails else None,
        "phone": phones[0] if phones else None,
        "extension": None,
        "organization": unit,
        "unit_type": _infer_unit_type(unit),
        "needs_review": True,
    }
    return parsed, [("warning", "free_text_person", "พบรายชื่อในข้อความอิสระ โปรดตรวจสอบการแยกข้อมูล")]


def parse_workbook(path: str | Path, original_name: str | None = None) -> list[ParsedSheet]:
    path = Path(path)
    workbook = load_workbook(path, read_only=False, data_only=True)
    parsed_sheets: list[ParsedSheet] = []
    try:
        for worksheet in workbook.worksheets:
            raw_rows: list[tuple[int, list[str]]] = []
            for row_number, row in enumerate(worksheet.iter_rows(values_only=True), start=1):
                values = [clean_text(value) for value in row]
                while values and not values[-1]:
                    values.pop()
                if any(values):
                    raw_rows.append((row_number, values))

            leading_text = " ".join(" ".join(values) for _, values in raw_rows[:20])
            building = _find_building(original_name or path.name, worksheet.title, leading_text)
            floor = _find_floor(worksheet.title, leading_text)
            active_header: dict[str, int] | None = None
            current_unit: str | None = worksheet.title
            rows: list[ParsedRow] = []

            for row_number, values in raw_rows:
                header = _header_map(values)
                if header:
                    active_header = header
                    rows.append(ParsedRow(row_number, "header", values))
                    continue

                parsed: dict[str, Any] | None = None
                issues: list[tuple[str, str, str]] = []
                if active_header:
                    def value_for(key: str) -> str:
                        index = active_header.get(key, -1)
                        return values[index] if 0 <= index < len(values) else ""

                    full_name = value_for("full_name")
                    email_match = EMAIL_RE.search(value_for("email"))
                    phone = value_for("phone")
                    extension = value_for("extension")
                    if full_name and not PHONE_RE.search(full_name) and (
                        _looks_like_person_name(full_name) or email_match or phone or extension or value_for("position")
                    ):
                        parsed = {
                            "full_name": full_name,
                            "nickname": value_for("nickname") or None,
                            "position": value_for("position") or None,
                            "email": email_match.group(0).lower() if email_match else None,
                            "phone": phone or None,
                            "extension": extension or None,
                            "organization": value_for("organization") or current_unit,
                            "unit_type": _infer_unit_type(value_for("organization") or current_unit),
                            "needs_review": not _looks_like_person_name(full_name),
                        }
                        if parsed["needs_review"]:
                            issues.append(("warning", "uncertain_name", "รูปแบบชื่อไม่ชัดเจน โปรดตรวจสอบ"))

                if parsed:
                    rows.append(ParsedRow(row_number, "employee_review" if issues else "employee", values, parsed, issues))
                    continue

                joined = " | ".join(value for value in values if value)
                free_parsed, free_issues = _extract_free_text(joined, current_unit)
                if free_parsed:
                    rows.append(ParsedRow(row_number, "employee_review", values, free_parsed, free_issues))
                    continue

                nonempty = [value for value in values if value]
                looks_section = len(nonempty) <= 2 and any(
                    word in joined for word in ("ฝ่าย", "แผนก", "สำนัก", "งาน", "บริษัท", "บริการ", "โครงการ", "ชั้น")
                )
                if looks_section:
                    current_unit = nonempty[0]
                    rows.append(ParsedRow(row_number, "section", values))
                else:
                    rows.append(ParsedRow(row_number, "note", values))

            parsed_sheets.append(ParsedSheet(worksheet.title, building, floor, rows))
    finally:
        workbook.close()
    return parsed_sheets
